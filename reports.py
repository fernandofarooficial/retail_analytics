"""Geração de relatórios em PDF/Excel (hoje só 'Identificados', ver routes/relatorios.py
e routes/mobile.py). PDF via reportlab, Excel via openpyxl — nenhum dos dois exige
dependência de sistema (diferente de weasyprint/wkhtmltopdf), então roda sem instalação
extra no VPS além de `pip install -r requirements.txt`."""

import io
import requests
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                 Spacer, Image, KeepTogether, HRFlowable)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (chave em identificados_lista, rótulo de exibição)
CAMPOS_IDENTIFICADOS = [
    ('person_id',           'ID'),
    ('full_name',           'Nome completo'),
    ('nickname',            'Apelido'),
    ('document',            'Documento'),
    ('birth_date',          'Data nasc.'),
    ('age',                 'Idade'),
    ('gender_name',         'Gênero'),
    ('person_type_name',    'Tipo'),
    ('notes',               'Observações'),
    ('phone',               'Telefone'),
    ('email',               'E-mail'),
    ('reviewed_by_name',    'Revisado por'),
    ('reviewed_at',         'Revisado em'),
    ('created_at',          'Criado em'),
    ('updated_at',          'Atualizado em'),
]

def _fmt(key, value):
    if value is None:
        return ''
    if key == 'birth_date':
        return value.strftime('%d/%m/%Y')
    if key in ('created_at', 'updated_at', 'reviewed_at'):
        return value.strftime('%d/%m/%Y %H:%M')
    return str(value)


def _fetch_photo(foto_url, timeout=5):
    if not foto_url:
        return None
    try:
        resp = requests.get(foto_url, timeout=timeout)
        resp.raise_for_status()
        return resp.content
    except Exception:
        return None


# ── PDF ──────────────────────────────────────────────────────────────────────

_styles = getSampleStyleSheet()
_cell_style = ParagraphStyle('cell', parent=_styles['Normal'], fontSize=7, leading=8.5)
_header_style = ParagraphStyle('cellHeader', parent=_styles['Normal'], fontSize=7,
                                leading=8.5, textColor=colors.white, fontName='Helvetica-Bold')
_titulo_style = ParagraphStyle('titulo', parent=_styles['Heading1'], fontSize=16, spaceAfter=2)
_sub_style = ParagraphStyle('sub', parent=_styles['Normal'], fontSize=9,
                             textColor=colors.HexColor('#6B7280'), spaceAfter=10)
_card_label_style = ParagraphStyle('cardLabel', parent=_styles['Normal'], fontSize=7,
                                    textColor=colors.HexColor('#6B7280'))
_card_value_style = ParagraphStyle('cardValue', parent=_styles['Normal'], fontSize=9,
                                    leading=11, spaceAfter=4)


def _cabecalho_flowables(loja_nome, data_ini, data_fim, total):
    periodo = f"Última atualização: {data_ini.strftime('%d/%m/%Y')} a {data_fim.strftime('%d/%m/%Y')}"
    return [
        Paragraph('Identificados', _titulo_style),
        Paragraph(
            f"{loja_nome or '—'} &nbsp;·&nbsp; {periodo} &nbsp;·&nbsp; {total} pessoa(s)",
            _sub_style),
    ]


def gerar_pdf_identificados_sem_foto(rows, loja_nome, data_ini, data_fim):
    """PDF em tabela (A3 paisagem), uma linha por pessoa, todas as colunas de people."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A3),
                             topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                             leftMargin=1 * cm, rightMargin=1 * cm)
    flow = _cabecalho_flowables(loja_nome, data_ini, data_fim, len(rows))

    header = [Paragraph(label, _header_style) for _, label in CAMPOS_IDENTIFICADOS]
    data = [header]
    for r in rows:
        data.append([
            Paragraph(_fmt(key, r.get(key)) or '—', _cell_style)
            for key, _ in CAMPOS_IDENTIFICADOS
        ])

    usable_width = landscape(A3)[0] - 2 * cm
    # nome/observações/e-mail ganham mais espaço; id/idade/data menos
    largura_relativa = {
        'person_id': 0.4, 'full_name': 1.6, 'nickname': 0.9, 'document': 0.9,
        'birth_date': 0.8, 'age': 0.5, 'gender_name': 0.7,
        'person_type_name': 0.8, 'notes': 1.6,
        'phone': 0.9, 'email': 1.3, 'reviewed_by_name': 1.0,
        'reviewed_at': 1.0, 'created_at': 1.0, 'updated_at': 1.0,
    }
    pesos = [largura_relativa.get(key, 1.0) for key, _ in CAMPOS_IDENTIFICADOS]
    total_peso = sum(pesos)
    col_widths = [usable_width * p / total_peso for p in pesos]

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0057A8')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D1D5DB')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    flow.append(table)
    doc.build(flow)
    return buf.getvalue()


def gerar_pdf_identificados_com_foto(rows, loja_nome, data_ini, data_fim):
    """PDF em cartões (A4 retrato) — foto + todos os campos por pessoa."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    flow = _cabecalho_flowables(loja_nome, data_ini, data_fim, len(rows))

    foto_size = 3.2 * cm
    campos_ficha = [(k, l) for k, l in CAMPOS_IDENTIFICADOS if k != 'full_name']

    for r in rows:
        photo_bytes = _fetch_photo(r.get('foto_url'))
        if photo_bytes:
            img_flow = Image(io.BytesIO(photo_bytes), width=foto_size, height=foto_size)
        else:
            ph = Table([['Sem foto']], colWidths=[foto_size], rowHeights=[foto_size])
            ph.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#E5E7EB')),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#6B7280')),
            ]))
            img_flow = ph

        nome = r.get('full_name') or f"Pessoa {r.get('person_id')}"
        campo_pares = []
        for key, label in campos_ficha:
            valor = _fmt(key, r.get(key)) or '—'
            campo_pares.append(Paragraph(f"<b>{label}:</b> {valor}", _card_value_style))
        meio = (len(campo_pares) + 1) // 2
        col_a = campo_pares[:meio]
        col_b = campo_pares[meio:]
        while len(col_b) < len(col_a):
            col_b.append(Spacer(1, 1))

        ficha = Table(list(zip(col_a, col_b)),
                       colWidths=[(A4[0] - 3 * cm - foto_size - 0.4 * cm) / 2] * 2)
        ficha.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        nome_par = Paragraph(f"<b>{nome}</b>", ParagraphStyle(
            'cardNome', parent=_styles['Normal'], fontSize=12, spaceAfter=4))
        corpo = Table([[nome_par], [ficha]], colWidths=[A4[0] - 3 * cm - foto_size - 0.4 * cm])
        corpo.setStyle(TableStyle([
            ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        card = Table([[img_flow, corpo]],
                      colWidths=[foto_size + 0.4 * cm, A4[0] - 3 * cm - foto_size - 0.4 * cm])
        card.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        flow.append(KeepTogether([
            card,
            HRFlowable(width='100%', thickness=0.5, color=colors.HexColor('#E5E7EB'),
                       spaceBefore=6, spaceAfter=10),
        ]))

    doc.build(flow)
    return buf.getvalue()


# ── Excel ────────────────────────────────────────────────────────────────────

def gerar_excel_identificados(rows, loja_nome, data_ini, data_fim):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Identificados'

    ws.append([f"Identificados — {loja_nome or '—'} — "
               f"atualizados de {data_ini.strftime('%d/%m/%Y')} a {data_fim.strftime('%d/%m/%Y')} "
               f"({len(rows)} pessoa(s))"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CAMPOS_IDENTIFICADOS))
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append([])

    header_row = 3
    for col, (_, label) in enumerate(CAMPOS_IDENTIFICADOS, start=1):
        cell = ws.cell(header_row, col, label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0057A8')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r_idx, r in enumerate(rows, start=header_row + 1):
        for col, (key, _) in enumerate(CAMPOS_IDENTIFICADOS, start=1):
            ws.cell(r_idx, col, _fmt(key, r.get(key)))

    for col, (key, label) in enumerate(CAMPOS_IDENTIFICADOS, start=1):
        largura = max(len(label), 12)
        if key in ('full_name', 'notes', 'email', 'reviewed_by_name'):
            largura = 28
        ws.column_dimensions[get_column_letter(col)].width = largura

    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
